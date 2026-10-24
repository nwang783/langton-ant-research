"""Build the continuation report with native Word mathematics and exact geometry."""
from pathlib import Path
import os
import subprocess
os.environ.setdefault('MPLCONFIGDIR', '/tmp/ant_report_20261022/matplotlib')
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OUT = ROOT / 'Langton_Ant_Research_Update_2026-10-22.docx'

fig, ax = plt.subplots(figsize=(6.0, 3.6))
colors = ('#A7B6BF', '#587B92', '#203F53')
for n, color in zip((1, 2, 3), colors):
    lo_x, hi_x = -2*n, 2*n+1
    lo_y, hi_y = -2*n+1, 2*n+2
    ax.add_patch(Rectangle((lo_x, lo_y), hi_x-lo_x, hi_y-lo_y,
                           fill=False, edgecolor=color, linewidth=1.3))
    points = [(x, y) for x in (lo_x, hi_x) for y in (lo_y, hi_y)]
    ax.scatter(*zip(*points), color=color, s=39, marker='s', zorder=4,
               label=f'n = {n}: four black cells')
ax.annotate('', xy=(0, 1.25), xytext=(0, 0),
            arrowprops={'arrowstyle': '-|>', 'color': '#B35F19', 'lw': 1.5})
ax.scatter([0], [0], s=25, color='#B35F19', zorder=5)
ax.annotate('head', (0, 0), xytext=(8, 7),
            textcoords='offset points', fontsize=8.5, color='#6D3B12')
ax.axhline(0, color='#D3D9DD', linewidth=.6, zorder=0)
ax.axvline(0, color='#D3D9DD', linewidth=.6, zorder=0)
ax.set_aspect('equal')
ax.set_xlim(-7, 8)
ax.set_ylim(-6, 9)
ax.set_xticks((-6, -4, -2, 0, 2, 4, 6, 8))
ax.set_yticks((-4, -2, 0, 2, 4, 6, 8))
ax.tick_params(labelsize=8, length=0)
ax.set_xlabel('x', fontsize=9)
ax.set_ylabel('y', fontsize=9, rotation=0)
for spine in ax.spines.values():
    spine.set_visible(False)
ax.legend(loc='center left', bbox_to_anchor=(1.05, .64), frameon=False, fontsize=9)
ax.text(1.05, .24, 'Endpoint tests pass.\nInfinite recurrence is\nexcluded by dynamics.',
        transform=ax.transAxes, fontsize=9, color='#333333', linespacing=1.5)
fig.subplots_adjust(left=.10, right=.60, bottom=.13, top=.98)
fig.savefig(HERE / 'endpoint_obstruction.png', dpi=230, bbox_inches='tight', facecolor='white')
plt.close(fig)

subprocess.run(['pandoc', str(HERE / 'report_body.md'), '-o', str(OUT),
                '--standalone', '--resource-path', str(HERE)], check=True)
doc = Document(OUT)
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.72)
section.bottom_margin = Inches(.72)
section.left_margin = Inches(.8)
section.right_margin = Inches(.8)
section.header_distance = Inches(.25)
section.footer_distance = Inches(.3)
for style in doc.styles:
    if style.type == 1:
        style.font.name = 'Calibri'
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        pf = style.paragraph_format
        pf.line_spacing = 1.05
        pf.space_after = Pt(6)
        if style.name.startswith('Heading'):
            style.font.bold = True
            pf.keep_with_next = True
            pf.space_before = Pt(15)
            pf.space_after = Pt(7)
styles = {style.name: style for style in doc.styles}
for name, size in [('Title', 25), ('Subtitle', 13), ('Heading 1', 15), ('Heading 2', 12.5)]:
    style = styles[name]
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if name == 'Title':
        style.font.bold = True
    if name in ('Title', 'Subtitle'):
        for border in style._element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
for name in ('Verbatim Char', 'Source Code'):
    if name in styles:
        styles[name].font.name = 'Consolas'
        styles[name].font.size = Pt(9)
paragraphs = doc.paragraphs
for i, paragraph in enumerate(paragraphs):
    if paragraph.style.name in ('Title', 'Subtitle', 'Date'):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for border in paragraph._p.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    if paragraph._p.xpath('.//m:oMath'):
        paragraph.paragraph_format.keep_together = True
    if i+1 < len(paragraphs) and paragraphs[i+1]._p.xpath('.//m:oMathPara'):
        paragraph.paragraph_format.keep_with_next = True
    if paragraph.style.name == 'Image Caption':
        paragraph.style = styles['Caption']
        for run in paragraph.runs:
            run.font.size = Pt(9.5)
    if paragraph._p.xpath('.//w:drawing'):
        paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        if paragraph.style.name.startswith(('Title', 'Subtitle', 'Heading')):
            run.font.color.rgb = RGBColor(0, 0, 0)

for index, table in enumerate(doc.tables):
    count = len(table.columns)
    widths = [2.3, 1.15, 1.15, 1.15, 1.15] if count == 5 else ([3.95, 2.95] if index == 0 else [2.65, 4.25])
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for j, column in enumerate(table.columns):
        column.width = Inches(widths[j])
    for row_index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, value in [('top', '85'), ('bottom', '85'), ('left', '105'), ('right', '105')]:
                node = OxmlElement('w:' + side)
                node.set(qn('w:w'), value)
                node.set(qn('w:type'), 'dxa')
                margins.append(node)
            tc_pr.append(margins)
            borders = OxmlElement('w:tcBorders')
            for side in ('top', 'left', 'bottom', 'right'):
                node = OxmlElement('w:' + side)
                node.set(qn('w:val'), 'single')
                node.set(qn('w:sz'), '4')
                node.set(qn('w:color'), 'D9D9D9')
                borders.append(node)
            tc_pr.append(borders)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if count == 5 and j else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0, 0, 0)
            if row_index == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E5ECF1')
                tc_pr.append(shade)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run()
run.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
run._r.addnext(field)
doc.core_properties.title = 'Langton Ant Research Continuation'
doc.core_properties.subject = 'Exact partial results for the finite-seed highway conjecture'
doc.core_properties.author = 'Research continuation prepared with ChatGPT'
doc.save(OUT)
print(OUT)
